"""TXT / DOCX / PDF / JSON renderers and the MURAILEX Forensic Evidence Package."""
from __future__ import annotations

import io
import json
import os
import zipfile
from datetime import datetime, timezone
from typing import Any

from ..canonical import canonical_json, sha256_hex
from ..pipeline import transcript as tx

TITLE = "MURAILEX FORENSIC VERBATIM TRANSCRIPT"
CONTROLLING = "The original audio recording is the controlling source."
FONT_CANDIDATES = [
    os.environ.get("MURAILEX_PDF_FONT", ""),
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
]
LATIN_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
]


def header_fields(ctx: dict[str, Any]) -> list[tuple[str, str]]:
    rec = ctx["recording"]
    rev = ctx["revision"]
    return [
        ("Recording", rec["title"]),
        ("Original file", rec["original_filename"]),
        ("Original SHA-256", rec["sha256"]),
        ("Original size (bytes)", str(rec["byte_size"])),
        ("MIME type", rec["mime_type"]),
        ("Duration", tx.fmt_ts(rec["duration_ms"] or 0)),
        ("Uploaded (UTC)", rec["uploaded_at"]),
        ("Revision", f"{rev['number']} ({rev['status']})"),
        ("Locked transcript SHA-256", rev["sha256"] or "NOT LOCKED"),
        ("Locked (UTC)", rev["locked_at"] or "—"),
        ("Engines", ", ".join(f"{e['provider']} {e['model']}" for e in ctx["engines"]) or "—"),
    ]


def render_txt(ctx: dict[str, Any], translation: dict[str, Any] | None = None) -> bytes:
    lines = [TITLE, CONTROLLING, ""]
    lines += [f"{k}: {v}" for k, v in header_fields(ctx)]
    if translation:
        lines += [f"Derived translation: {translation['mode']} ({translation['provider']} {translation['model']})",
                  f"Translation SHA-256: {translation['sha256']}",
                  "This translation is a separate derived document. The source-language forensic transcript controls."]
    lines += ["", "-" * 72, ""]
    if translation:
        for s in translation["segments"]:
            lines.append(f"[{tx.fmt_ts(s['start_ms'])}] {s['speaker_label']}")
            if translation["mode"] == "bilingual":
                lines.append(f"⁧{s['source_text']}⁩")
            lines.append(s["translation"])
            lines.append("")
    else:
        for ts, label, text in tx.plain_lines(ctx["content"]):
            lines.append(f"[{ts}] ⁧{label}⁩")
            lines.append(f"⁨{text}⁩")
            lines.append("")
    return ("\n".join(lines) + "\n").encode("utf-8")


def render_json(ctx: dict[str, Any]) -> bytes:
    doc = {
        "title": TITLE,
        "controlling_source": CONTROLLING,
        "recording": ctx["recording"],
        "revision": ctx["revision"],
        "engines": ctx["engines"],
        "transcript": ctx["content"],
    }
    return json.dumps(doc, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")


def _rtl_paragraph(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)


def _rtl_run(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rpr.append(rtl)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:cs"), "Noto Naskh Arabic")


def _is_rtl(text: str) -> bool:
    from ..pipeline.text import ARABIC_CHAR, LATIN_CHAR

    ar = len(ARABIC_CHAR.findall(text))
    return ar > 0 and ar >= len(LATIN_CHAR.findall(text))


def render_docx(ctx: dict[str, Any], translation: dict[str, Any] | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = TITLE
    h = doc.add_heading(TITLE, level=1)
    h.alignment = 1  # type: ignore[assignment]
    p = doc.add_paragraph()
    p.add_run(CONTROLLING).bold = True
    table = doc.add_table(rows=0, cols=2)
    for k, v in header_fields(ctx):
        row = table.add_row().cells
        row[0].text = k
        row[1].text = str(v)
    if translation:
        doc.add_paragraph(
            f"Derived translation ({translation['mode']}, {translation['provider']} {translation['model']}). "
            f"SHA-256 {translation['sha256']}. The source-language forensic transcript controls."
        )
    doc.add_paragraph("")
    entries: list[tuple[str, str, list[str]]] = []
    if translation:
        for s in translation["segments"]:
            texts = [s["source_text"], s["translation"]] if translation["mode"] == "bilingual" else [s["translation"]]
            entries.append((tx.fmt_ts(s["start_ms"]), s["speaker_label"], texts))
    else:
        for ts, label, text in tx.plain_lines(ctx["content"]):
            entries.append((ts, label, [text]))
    for ts, label, texts in entries:
        head = doc.add_paragraph()
        r = head.add_run(f"[{ts}]  ")
        r.font.size = Pt(9)
        lr = head.add_run(label)
        lr.bold = True
        if _is_rtl(label):
            _rtl_run(lr)
        for text in texts:
            para = doc.add_paragraph()
            run = para.add_run(text)
            run.font.size = Pt(12)
            if _is_rtl(text):
                _rtl_paragraph(para)
                para.alignment = 2  # type: ignore[assignment]
                _rtl_run(run)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_fonts_registered = False


def _register_fonts() -> tuple[str, str]:
    global _fonts_registered
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    ar = next((f for f in FONT_CANDIDATES if f and os.path.exists(f)), None)
    la = next((f for f in LATIN_FONT_CANDIDATES if os.path.exists(f)), None)
    if not ar or not la:
        raise RuntimeError("Arabic/Latin PDF fonts are not installed (fonts-noto-core, fonts-dejavu-core).")
    if not _fonts_registered:
        pdfmetrics.registerFont(TTFont("MXArabic", ar))
        pdfmetrics.registerFont(TTFont("MXLatin", la))
        _fonts_registered = True
    return "MXArabic", "MXLatin"


def _visual(text: str) -> str:
    """Shape Arabic and reorder bidirectional text for PDF glyph placement (display only)."""
    import arabic_reshaper
    from bidi.algorithm import get_display

    return get_display(arabic_reshaper.reshape(text))


def _font_runs(text: str, ar_font: str, la_font: str, size: int | float) -> str:
    from xml.sax.saxutils import escape

    from ..pipeline.text import ARABIC_CHAR

    out = []
    cur_font = None
    buf = ""
    for ch in text:
        f = ar_font if (ARABIC_CHAR.match(ch) or "ﭐ" <= ch <= "﻿") else (la_font if ch.strip() else cur_font or la_font)
        if f != cur_font and buf:
            out.append(f'<font name="{cur_font}" size="{size}">{escape(buf)}</font>')
            buf = ""
        cur_font = f
        buf += ch
    if buf:
        out.append(f'<font name="{cur_font}" size="{size}">{escape(buf)}</font>')
    return "".join(out)


def render_pdf(ctx: dict[str, Any], translation: dict[str, Any] | None = None) -> bytes:
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    ar_font, la_font = _register_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=TITLE, author="MURAILEX")
    ltr = ParagraphStyle("ltr", fontName=la_font, fontSize=10, leading=14, alignment=TA_LEFT)
    rtl = ParagraphStyle("rtl", fontName=ar_font, fontSize=12, leading=19, alignment=TA_RIGHT)
    title = ParagraphStyle("title", fontName=la_font, fontSize=15, leading=20, alignment=1, spaceAfter=6)
    story: list[Any] = [Paragraph(TITLE, title), Paragraph(f"<b>{CONTROLLING}</b>", ltr), Spacer(1, 4 * mm)]
    rows = [[Paragraph(k, ltr), Paragraph(_font_runs(str(v), ar_font, la_font, 9), ltr)] for k, v in header_fields(ctx)]
    t = Table(rows, colWidths=[45 * mm, 125 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, "#9999aa"), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Spacer(1, 5 * mm)]
    if translation:
        story.append(Paragraph(
            f"Derived translation ({translation['mode']}, {translation['provider']} {translation['model']}), "
            f"SHA-256 {translation['sha256']}. The source-language forensic transcript controls.", ltr))
        story.append(Spacer(1, 3 * mm))
    entries: list[tuple[str, str, list[str]]] = []
    if translation:
        for s in translation["segments"]:
            texts = [s["source_text"], s["translation"]] if translation["mode"] == "bilingual" else [s["translation"]]
            entries.append((tx.fmt_ts(s["start_ms"]), s["speaker_label"], texts))
    else:
        for ts, label, text in tx.plain_lines(ctx["content"]):
            entries.append((ts, label, [text]))
    for ts, label, texts in entries:
        story.append(Paragraph(f"[{ts}]  " + _font_runs(_visual(label), ar_font, la_font, 10), ltr))
        for text in texts:
            style = rtl if _is_rtl(text) else ltr
            story.append(Paragraph(_font_runs(_visual(text), ar_font, la_font, 12 if style is rtl else 11), style))
        story.append(Spacer(1, 2.5 * mm))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont(la_font, 7.5)
        canvas.drawString(18 * mm, 10 * mm, f"{TITLE} · {CONTROLLING} · Transcript SHA-256 {ctx['revision']['sha256'] or 'NOT LOCKED'}")
        canvas.drawRightString(A4[0] - 18 * mm, 6 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def build_package(ctx: dict[str, Any], extras: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
    """Return (zip bytes, manifest). Every file is listed in SHA256SUMS.txt."""
    files: dict[str, bytes] = {
        "transcript/transcript.txt": render_txt(ctx),
        "transcript/transcript.docx": render_docx(ctx),
        "transcript/transcript.pdf": render_pdf(ctx),
        "transcript/transcript.json": render_json(ctx),
        "revision_history.json": json.dumps(extras["revisions"], ensure_ascii=False, indent=2).encode(),
        "audit_log.json": json.dumps(extras["audit"], ensure_ascii=False, indent=2).encode(),
        "provider_runs.json": json.dumps(extras["provider_runs"], ensure_ascii=False, indent=2).encode(),
        "disputes.json": json.dumps(extras["disputes"], ensure_ascii=False, indent=2).encode(),
    }
    for tr in extras.get("translations", []):
        base = f"translations/{tr['mode']}-{tr['id'][:8]}"
        files[f"{base}.txt"] = render_txt(ctx, tr)
        files[f"{base}.docx"] = render_docx(ctx, tr)
        files[f"{base}.pdf"] = render_pdf(ctx, tr)
        files[f"{base}.json"] = json.dumps(tr, ensure_ascii=False, indent=2).encode()
    manifest = {
        "title": TITLE,
        "controlling_source": CONTROLLING,
        "package": "MURAILEX Forensic Evidence Package",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "generated_by": extras["actor"],
        "recording": ctx["recording"],
        "original_sha256": ctx["recording"]["sha256"],
        "original_sha256_reverified_at_export": extras["reverified"],
        "locked_transcript": {"revision_id": ctx["revision"]["id"], "number": ctx["revision"]["number"],
                              "sha256": ctx["revision"]["sha256"], "locked_at": ctx["revision"]["locked_at"],
                              "locked_by": ctx["revision"]["locked_by"]},
        "engines": ctx["engines"],
        "method": ctx["content"].get("method"),
        "audit_chain": extras["audit_chain"],
        "translations": [{"id": t["id"], "mode": t["mode"], "sha256": t["sha256"]} for t in extras.get("translations", [])],
        "files": {name: {"sha256": sha256_hex(data), "bytes": len(data)} for name, data in sorted(files.items())},
        "notice": "Derived documents. The original audio recording is the controlling source.",
    }
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode()
    sums = "".join(f"{sha256_hex(data)}  {name}\n" for name, data in sorted(files.items()))
    files["SHA256SUMS.txt"] = sums.encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, data in sorted(files.items()):
            info = zipfile.ZipInfo(f"MURAILEX Forensic Evidence Package/{name}", date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(info, data)
    return buf.getvalue(), manifest


def content_hash(recording_sha: str, number: int, parent_sha: str | None, content: dict[str, Any]) -> str:
    return sha256_hex(canonical_json({"original_sha256": recording_sha, "revision": number, "parent_sha256": parent_sha, "content": content}))